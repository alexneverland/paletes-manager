from contextlib import closing
from datetime import date
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import sqlite3
import tkinter as tk

import pytest

from paletes import db
from paletes import export
from paletes.constants import CARRIERS
from paletes.ui import PaletesApp


@pytest.fixture
def database(tmp_path, monkeypatch):
    path = tmp_path / 'data' / 'paletes.db'
    monkeypatch.setattr(db, 'DB_PATH', str(path))
    monkeypatch.setattr(db, 'BACKUP_DIR', str(path.parent / 'backups'))
    # A dialog would block unattended execution when reproducing the old bug.
    monkeypatch.setattr(db.messagebox, 'showerror', lambda *args: None)
    return path


def test_first_start_creates_parent_and_usable_tables(database):
    db.init_database()
    assert database.exists()
    assert db.fetch_entries(date.today().isoformat()) == []
    assert db.fetch_predictions(date.today().isoformat()) == []


def test_shared_database_uses_rollback_journal(database):
    database.parent.mkdir()
    db.init_database()
    with closing(db.get_connection()) as conn:
        assert conn.execute('PRAGMA journal_mode').fetchone()[0] == 'delete'
        assert conn.execute('PRAGMA synchronous').fetchone()[0] == 2


def test_backup_includes_committed_wal_rows(database):
    database.parent.mkdir()
    db.init_database()
    with closing(sqlite3.connect(database)) as writer:
        writer.execute('PRAGMA journal_mode=WAL')
        writer.execute('PRAGMA wal_autocheckpoint=0')
        writer.execute("INSERT INTO entries(entry_date, carrier, code) VALUES (?, ?, ?)",
                       (date.today().isoformat(), CARRIERS[0], 'P1'))
        writer.commit()
        db.auto_backup_db()
        backup = next(Path(db.BACKUP_DIR).glob('*.db'))
        with closing(sqlite3.connect(backup)) as restored:
            assert restored.execute('SELECT code FROM entries').fetchall() == [('P1',)]
            assert restored.execute('PRAGMA integrity_check').fetchone() == ('ok',)


def test_old_wal_database_is_migrated_without_losing_rows(database):
    database.parent.mkdir()
    with closing(sqlite3.connect(database)) as conn:
        conn.execute('PRAGMA journal_mode=WAL')
        conn.execute('CREATE TABLE preserved (value TEXT)')
        conn.execute("INSERT INTO preserved VALUES ('existing')")
        conn.commit()
    assert db.init_database()
    with closing(db.get_connection()) as conn:
        assert conn.execute('SELECT value FROM preserved').fetchall() == [('existing',)]
        assert conn.execute('PRAGMA journal_mode').fetchone() == ('delete',)


def test_two_local_clients_do_not_lose_concurrent_inserts(database):
    db.init_database()

    def insert_batch(prefix):
        for index in range(20):
            db.insert_entry(date.today().isoformat(), CARRIERS[0], f'{prefix}-{index}',
                            'Alice', '1', 'ΟΧΙ', '', '')

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(insert_batch, ['PC1', 'PC2']))
    rows = db.fetch_entries(date.today().isoformat())
    assert len(rows) == 40
    assert len({row[2] for row in rows}) == 40


def test_failed_backup_is_not_published(database, monkeypatch):
    db.init_database()

    def fail_publish(*args):
        raise OSError('simulated unavailable share')

    monkeypatch.setattr(db.os, 'replace', fail_publish)
    assert db.auto_backup_db() is False
    assert list(Path(db.BACKUP_DIR).iterdir()) == []


def test_excel_export_contains_stored_data(database, monkeypatch):
    import openpyxl

    db.init_database()
    db.insert_entry(date.today().isoformat(), CARRIERS[0], 'P1', 'Αθήνα', '2', 'ΟΧΙ', '5', 'test')
    desktop = database.parent / 'Desktop'
    desktop.mkdir()
    monkeypatch.setenv('USERPROFILE', str(database.parent))
    monkeypatch.setattr(export.webbrowser, 'open', lambda url: True)
    path, error = export.export_dataframe_to_excel(db.fetch_main_export(date.today().isoformat()), 'test')
    assert error is None
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        assert list(workbook.active.values)[1] == (CARRIERS[0], 'P1', 'Αθήνα', '2', 'ΟΧΙ', '5', 'test')
    finally:
        workbook.close()


@pytest.fixture(scope='module')
def tk_root():
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()


@pytest.fixture
def app(database, tk_root):
    # mkdir isolates the UI regressions from the independent first-start bug.
    database.parent.mkdir()
    root = tk.Toplevel(tk_root)
    root.withdraw()
    instance = PaletesApp(root)
    yield instance
    root.destroy()


@pytest.mark.parametrize('mode', ['main', 'prediction'])
def test_refresh_preserves_unsaved_edit(app, mode):
    carrier = CARRIERS[0]
    today = date.today().isoformat()
    if mode == 'main':
        db.insert_entry(today, carrier, 'P1', 'Alice', '1', 'ΟΧΙ', '', '')
    else:
        db.insert_prediction(today, carrier, 'Alice', 'Παλέτα', 1, '')
    app.set_mode(mode)
    table = app.tables[carrier]
    table.selection_set(table.get_children()[0])
    if mode == 'main':
        app.enter_main_edit_mode(carrier)
        field = app.main_widgets[carrier]['entries'][0]
    else:
        app.enter_prediction_edit_mode(carrier)
        field = app.prediction_widgets[carrier]['entries'][1]
    field.delete(0, tk.END)
    field.insert(0, '123')
    app.auto_refresh()
    assert app.edit_mode
    assert field.get() == '123'
    app.refresh_data()
    assert app.edit_mode
    assert field.get() == '123'


def test_column_sort_is_not_overridden_by_name_grouping(app):
    carrier = CARRIERS[0]
    for code, name, invoice in [('P1', 'Alice', '1'), ('P2', 'Bob', '2'), ('P3', 'Alice', '3')]:
        db.insert_entry(date.today().isoformat(), carrier, code, name, invoice, 'ΟΧΙ', '', '')
    app.load_main_data(date.today())
    app.sort_column_handler(carrier, 'invoice')
    table = app.tables[carrier]
    assert [str(table.item(iid, 'values')[2]) for iid in table.get_children()] == ['1', '2', '3']
    app.sort_column_handler(carrier, 'invoice')
    assert [str(table.item(iid, 'values')[2]) for iid in table.get_children()] == ['3', '2', '1']
